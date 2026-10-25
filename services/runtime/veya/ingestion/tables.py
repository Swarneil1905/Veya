"""Table values are data only: formulas are never executed or exported as formulas."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from veya.domain.models import DomainError


def extract_table(kind: str, content: bytes) -> dict:
    blocks: list[dict] = []
    issues: list[dict] = []
    book = None
    if kind == 'csv':
        sheets: Any = [('CSV', csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True))]
    else:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) > 20:
            book.close()
            raise DomainError('too_large', 'Workbook exceeds 20 sheets')
        sheets = [(sheet.title, sheet.values) for sheet in book.worksheets]
    try:
        for sheet, source in sheets:
            rows = iter(source)
            headers = [str(v or '').strip().lower() for v in next(rows, [])]
            if not headers or len(headers) > 100:
                raise DomainError('invalid_input', 'Table needs 1 to 100 columns')
            if len(set(headers)) != len(headers) or '' in headers:
                issues.append({'code': 'ambiguous_headers', 'locator': sheet + '!1'})
            for row_number, row in enumerate(rows, 2):
                if row_number > 2001 or len(blocks) + len(row) > 10000:
                    raise DomainError('too_large', 'Table exceeds 2000 rows or 10000 cells')
                if len(row) != len(headers):
                    issues.append({'code': 'row_width', 'locator': f'{sheet}!{row_number}'})
                for column, value in enumerate(row):
                    locator = f'{sheet}!{get_column_letter(column + 1)}{row_number}'
                    text = (
                        value.isoformat()
                        if isinstance(value, (date, datetime))
                        else str(value if value is not None else '')
                    )
                    if len(text) > 4000:
                        raise DomainError('too_large', 'Cell exceeds 4000 characters')
                    header = headers[column] if column < len(headers) else ''
                    typed, value_type = text, 'text'
                    if text.lstrip().startswith(('=', '+', '@')) or (
                        text.lstrip().startswith('-') and not re.fullmatch(r'-\d+(?:\.\d+)?', text)
                    ):
                        value_type = 'untrusted_text'
                        issues.append({'code': 'formula', 'locator': locator})
                    elif header in {'amount', 'quantity', 'price', 'total'}:
                        try:
                            number = Decimal(text)
                            if not number.is_finite() or not re.fullmatch(r'-?\d+(?:\.\d+)?', text):
                                raise InvalidOperation()
                            value_type = 'decimal'
                        except InvalidOperation:
                            issues.append({'code': 'invalid_decimal', 'locator': locator})
                    elif header in {'date', 'due_date', 'invoice_date'}:
                        try:
                            typed = date.fromisoformat(text[:10]).isoformat()
                            if len(text) != 10 and not isinstance(value, (date, datetime)):
                                raise ValueError()
                            value_type = 'date'
                        except ValueError:
                            issues.append({'code': 'invalid_date', 'locator': locator})
                    elif header == 'currency':
                        if re.fullmatch('[A-Z]{3}', text):
                            value_type = 'currency'
                        else:
                            issues.append({'code': 'invalid_currency', 'locator': locator})
                    blocks.append(
                        {
                            'locator': locator,
                            'text': text,
                            'value': typed,
                            'value_type': value_type,
                            'field': header,
                        }
                    )
        if not blocks:
            issues.append({'code': 'empty_table', 'locator': ''})
        return {'blocks': blocks, 'issues': issues}
    finally:
        if book is not None:
            book.close()
